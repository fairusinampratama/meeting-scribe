"""Render a summary markdown file to HTML + DOCX.

No pandoc / node / LibreOffice required: uses `markdown` + `python-docx`, which
install cleanly on a locked-down machine without admin rights.
"""
import os
import re

INK, MUTED, ACC, TS = (0x1C,0x1B,0x19), (0x6B,0x68,0x62), (0x8A,0x5A,0x2B), (0x2F,0x6B,0x57)

CSS = """
:root{--bg:#faf9f7;--panel:#fff;--ink:#1c1b19;--muted:#6b6862;--line:#e3e0da;
--accent:#8a5a2b;--accent-soft:#f4ece1;--ts:#2f6b57;--ts-soft:#e9f2ee;}
@media (prefers-color-scheme:dark){:root:not([data-theme="light"]){
--bg:#16161a;--panel:#1d1d22;--ink:#e8e6e1;--muted:#9a968e;--line:#2f2f36;
--accent:#d8a05f;--accent-soft:#2a2119;--ts:#7fc4a6;--ts-soft:#1a2a24;}}
:root[data-theme="dark"]{--bg:#16161a;--panel:#1d1d22;--ink:#e8e6e1;--muted:#9a968e;
--line:#2f2f36;--accent:#d8a05f;--accent-soft:#2a2119;--ts:#7fc4a6;--ts-soft:#1a2a24;}
*{box-sizing:border-box}
body{background:var(--bg);color:var(--ink);margin:0;
font:16px/1.65 ui-sans-serif,-apple-system,"Segoe UI",Roboto,Helvetica,Arial,sans-serif}
.wrap{max-width:860px;margin:0 auto;padding:48px 28px 96px}
h1{font-size:1.9rem;line-height:1.25;margin:0 0 .2em;letter-spacing:-.02em}
h2{font-size:1.28rem;margin:2.4em 0 .6em;padding-bottom:.3em;border-bottom:2px solid var(--line)}
h3{font-size:1.02rem;margin:1.8em 0 .5em;color:var(--accent)}
hr{border:0;border-top:1px solid var(--line);margin:2.4em 0}
a{color:var(--accent)}
code{font:13px/1.5 ui-monospace,"Cascadia Mono",Consolas,monospace;
background:var(--accent-soft);padding:1px 5px;border-radius:4px}
code.ts{background:var(--ts-soft);color:var(--ts);font-weight:600;white-space:nowrap}
table{border-collapse:collapse;width:100%;margin:1.1em 0;font-size:.93rem;
background:var(--panel);border:1px solid var(--line);border-radius:8px;overflow:hidden}
th{background:var(--accent-soft);text-align:left;font-weight:650;font-size:.82rem;
letter-spacing:.04em;text-transform:uppercase;color:var(--muted)}
th,td{padding:9px 12px;border-bottom:1px solid var(--line);vertical-align:top}
tr:last-child td{border-bottom:0}
ol,ul{padding-left:1.3em} li{margin:.32em 0}
strong{font-weight:650} em{color:var(--muted)}
.tablescroll{overflow-x:auto}
h1 + p{color:var(--muted);font-size:.92rem;border-bottom:1px solid var(--line);padding-bottom:1.2em}
"""


def to_html(src, dst, title):
    import markdown
    body = markdown.markdown(open(src, encoding="utf-8").read(),
                             extensions=["tables", "sane_lists", "attr_list"])
    body = re.sub(r"<code>\[(\d{1,3}:\d{2}(?:[\u2013-]\d{1,3}:\d{2})?)\]</code>",
                  r'<code class="ts">[\1]</code>', body)
    html = (f'<!doctype html><html><head><meta charset="utf-8">'
            f'<meta name="viewport" content="width=device-width,initial-scale=1">'
            f'<title>{title}</title><style>{CSS}</style></head><body><div class="wrap">\n'
            f'{body}\n</div><script>\n'
            "document.querySelectorAll('table').forEach(function(t){"
            "var d=document.createElement('div');d.className='tablescroll';"
            "t.parentNode.insertBefore(d,t);d.appendChild(t);});\n"
            "</script></body></html>")
    open(dst, "w", encoding="utf-8").write(html)
    return len(html)


def to_docx(src, dst):
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    ink, muted, acc, ts = (RGBColor(*INK), RGBColor(*MUTED), RGBColor(*ACC), RGBColor(*TS))
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)      # A4
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin = sec.bottom_margin = Cm(2.0)

    st = doc.styles["Normal"]
    st.font.name = "Calibri"; st.font.size = Pt(10.5); st.font.color.rgb = ink
    st.paragraph_format.space_after = Pt(6); st.paragraph_format.line_spacing = 1.15
    for name, size, color, before, after in [("Heading 1",20,ink,0,8),
                                             ("Heading 2",14,ink,18,6),
                                             ("Heading 3",11,acc,12,4)]:
        s = doc.styles[name]
        s.font.name = "Calibri"; s.font.size = Pt(size); s.font.bold = True
        s.font.color.rgb = color
        s.paragraph_format.space_before = Pt(before); s.paragraph_format.space_after = Pt(after)
        s.paragraph_format.keep_with_next = True

    def shade(cell, fill):
        tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd")
        shd.set(qn("w:val"),"clear"); shd.set(qn("w:color"),"auto"); shd.set(qn("w:fill"),fill)
        tcPr.append(shd)

    def bottom_border(p):
        pPr = p._p.get_or_add_pPr(); pbdr = OxmlElement("w:pBdr"); b = OxmlElement("w:bottom")
        b.set(qn("w:val"),"single"); b.set(qn("w:sz"),"6"); b.set(qn("w:space"),"4")
        b.set(qn("w:color"),"D8D4CC"); pbdr.append(b); pPr.append(pbdr)

    TOKEN = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+?`)")
    TSPAT = re.compile(r"^\[\d{1,3}:\d{2}(?:[\u2013-]\d{1,3}:\d{2})?\]$")

    def add_runs(p, text):
        for part in TOKEN.split(text):
            if not part: continue
            if part.startswith("**") and part.endswith("**"):
                p.add_run(part[2:-2]).bold = True
            elif part.startswith("`") and part.endswith("`"):
                inner = part[1:-1]; r = p.add_run(inner)
                r.font.name = "Consolas"; r.font.size = Pt(9)
                r.font.color.rgb = ts if TSPAT.match(inner) else acc
                if TSPAT.match(inner): r.bold = True
            elif part.startswith("*") and part.endswith("*") and len(part) > 2:
                r = p.add_run(part[1:-1]); r.italic = True; r.font.color.rgb = muted
            else:
                p.add_run(part)

    def split_row(line): return [c.strip() for c in line.strip().strip("|").split("|")]

    lines = open(src, encoding="utf-8").read().splitlines()
    i = 0; first = False; n_head = n_tab = 0
    while i < len(lines):
        s = lines[i].strip()
        if not s: i += 1; continue
        if s == "---":
            p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(2)
            bottom_border(p); i += 1; continue
        if s.startswith("|") and i+1 < len(lines) and re.match(r"^\|[\s:|-]+\|$", lines[i+1].strip()):
            header = split_row(s); i += 2; rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(split_row(lines[i])); i += 1
            t = doc.add_table(rows=1, cols=len(header)); t.style = "Table Grid"; t.autofit = True
            for j, h in enumerate(header):
                c = t.rows[0].cells[j]; c.text = ""
                p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(2)
                r = p.add_run(h.replace("**","").upper())
                r.bold = True; r.font.size = Pt(8); r.font.color.rgb = muted
                shade(c, "F4ECE1")
            for row in rows:
                cells = t.add_row().cells
                for j, val in enumerate(row[:len(header)]):
                    cells[j].text = ""; p = cells[j].paragraphs[0]
                    p.paragraph_format.space_after = Pt(2); add_runs(p, val)
                    for r in p.runs:
                        if r.font.size is None: r.font.size = Pt(9.5)
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
            n_tab += 1; continue
        m = re.match(r"^(#{1,3})\s+(.*)$", s)
        if m:
            lvl = len(m.group(1))
            p = doc.add_paragraph(style=f"Heading {lvl}"); add_runs(p, m.group(2))
            for r in p.runs: r.font.color.rgb = acc if lvl == 3 else ink; r.bold = True
            n_head += 1; i += 1; continue
        mb = re.match(r"^[-*]\s+(.*)$", s); mn = re.match(r"^(\d+)\.\s+(.*)$", s)
        if mb or mn:
            p = doc.add_paragraph(style="List Bullet" if mb else "List Number")
            p.paragraph_format.space_after = Pt(3)
            add_runs(p, mb.group(1) if mb else mn.group(2)); i += 1; continue
        p = doc.add_paragraph(); add_runs(p, s)
        if not first and i < 6:
            for r in p.runs: r.font.size = Pt(9.5); r.font.color.rgb = muted
            bottom_border(p); first = True
        i += 1

    fp = doc.sections[0].footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run()
    b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = " PAGE "
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end")
    run._r.append(b); run._r.append(it); run._r.append(e)
    run.font.size = Pt(8); run.font.color.rgb = muted

    doc.save(dst)
    return n_head, n_tab
